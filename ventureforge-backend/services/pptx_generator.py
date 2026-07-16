from __future__ import annotations

from io import BytesIO
from typing import Any

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE as MSO_SHAPE
from pptx.enum.shapes import MSO_CONNECTOR_TYPE as MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt


SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def _add_image(slide, image_bytes: bytes | None, left: float, top: float, width: float, height: float, *, opacity: float = 0.25):
    if not image_bytes:
        return None
    from io import BytesIO as _BIO
    stream = _BIO(image_bytes)
    pic = slide.shapes.add_picture(stream, Inches(left), Inches(top), Inches(width), Inches(height))
    # Send to back so text stays readable
    sp = pic._element
    sp.getparent().insert(0, sp)
    # Apply transparency via alpha modulation on the image fill
    if opacity < 1.0:
        from lxml import etree
        ns = "http://schemas.openxmlformats.org/drawingml/2006/main"
        blend_fill = pic._element.find(f".//{{{ns}}}blipFill")
        if blend_fill is not None:
            blip = blend_fill.find(f"{{{ns}}}blip")
            if blip is not None:
                alpha_mod = etree.SubElement(blip, f"{{{ns}}}alphaModFix")
                alpha_mod.set("amt", str(int(opacity * 100000)))
    return pic


def _rgb(hex_color: str) -> RGBColor:
    value = hex_color.strip().lstrip("#")
    return RGBColor.from_string(value)


def _configure_slide(slide, *, bg: str = "#0A0A0F"):
    pass # Let the template or default theme handle the background


def _add_textbox(
    slide,
    text: str,
    left: float,
    top: float,
    width: float,
    height: float,
    *,
    font_size: int = 18,
    bold: bool = False,
    color: str | None = None,
    align: PP_ALIGN = PP_ALIGN.LEFT,
    font_name: str = "Aptos",
) -> Any:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    r = p.runs[0]
    r.font.size = Pt(font_size)
    r.font.bold = bold
    r.font.name = font_name
    if color:
        r.font.color.rgb = _rgb(color)
    return box


def _add_card(slide, title: str, value: str, x: float, y: float, w: float, h: float, *, accent: str = "#6C47FF"):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.background() # Inherit background
    shape.line.color.rgb = _rgb(accent)
    shape.line.width = Pt(1.4)
    _add_textbox(slide, title, x + 0.16, y + 0.12, w - 0.32, 0.3, font_size=10, bold=True)
    _add_textbox(slide, value, x + 0.16, y + 0.42, w - 0.32, h - 0.5, font_size=18, bold=True)
    return shape


def _add_tag(slide, text: str, x: float, y: float, w: float, h: float, *, fill: str = "#0F172A", color: str = "#F0F0F0"):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(fill)
    shape.line.color.rgb = _rgb(fill)
    tf = shape.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = PP_ALIGN.CENTER
    r = p.runs[0]
    r.font.size = Pt(10)
    r.font.bold = True
    r.font.color.rgb = _rgb(color)
    return shape


def _safe_str(value: Any, fallback: str = "—") -> str:
    if value is None:
        return fallback
    if isinstance(value, (int, float)):
        return f"{value:,}"
    text = str(value).strip()
    return text if text else fallback


def _market_values(state: dict[str, Any]) -> tuple[str, str, str]:
    market = state.get("market") or {}
    return _safe_str(market.get("tam")), _safe_str(market.get("sam")), _safe_str(market.get("som"))


def _make_slide_title(slide, title: str, subtitle: str):
    _add_textbox(slide, title, 0.7, 0.45, 10.5, 0.6, font_size=26, bold=True)
    _add_textbox(slide, subtitle, 0.7, 0.98, 10.5, 0.35, font_size=11)
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.7), Inches(1.34), Inches(0.58), Inches(0.06))
    accent.fill.solid()
    accent.fill.fore_color.rgb = _rgb("#00D4AA")
    accent.line.fill.background()


def _split_business(state: dict[str, Any]) -> dict[str, Any]:
    return state.get("business_plan") or {}


def add_cover_slide(pptx: Presentation, state: dict[str, Any], images: dict[str, bytes] | None = None):
    slide = pptx.slides.add_slide(pptx.slide_layouts[6])
    _configure_slide(slide)
    _add_image(slide, (images or {}).get("cover"), 8.0, 0.5, 5.0, 6.5, opacity=0.18)
    startup_name = _safe_str(state.get("startup_name") or state.get("idea"), "Startup")
    tagline = _safe_str(((state.get("pitch_deck") or {}).get("brand") or {}).get("tagline"), "Investor-ready startup deck")
    _add_textbox(slide, startup_name, 0.72, 0.7, 8.8, 0.8, font_size=30, bold=True)
    _add_textbox(slide, tagline, 0.72, 1.5, 8.8, 0.4, font_size=16)
    _add_textbox(slide, "AI-generated startup narrative, market, architecture, and financials", 0.72, 1.95, 8.8, 0.35, font_size=11)
    market_tam, market_sam, market_som = _market_values(state)
    _add_card(slide, "TAM", market_tam, 0.72, 3.0, 2.35, 1.3, accent="#6C47FF")
    _add_card(slide, "SAM", market_sam, 3.2, 3.0, 2.35, 1.3, accent="#00D4AA")
    _add_card(slide, "SOM", market_som, 5.68, 3.0, 2.35, 1.3, accent="#4DA3FF")
    _add_tag(slide, "Powered by VentureForge", 0.72, 6.55, 2.3, 0.38, fill="#111118")


def add_executive_summary_slide(pptx: Presentation, state: dict[str, Any], images: dict[str, bytes] | None = None):
    slide = pptx.slides.add_slide(pptx.slide_layouts[6])
    _configure_slide(slide)
    _add_image(slide, (images or {}).get("executive_summary"), 9.5, 1.6, 3.5, 2.4, opacity=0.3)
    _make_slide_title(slide, "Executive Summary", "What the startup is, why it matters, and how it makes money")
    bp = _split_business(state)
    left = [
        ("Problem", _safe_str(bp.get("problem_statement"), "Problem statement pending")),
        ("Solution", _safe_str(bp.get("solution"), "Solution pending")),
        ("Value", _safe_str(bp.get("value_proposition"), "Value proposition pending")),
        ("Revenue", _safe_str(bp.get("revenue_model"), "Subscription / usage-based")),
    ]
    x = 0.72
    for title, value in left:
        _add_card(slide, title, value, x, 1.8, 2.95, 1.1, accent="#6C47FF" if title != "Solution" else "#00D4AA")
        x += 3.12
    _add_textbox(slide, "Target market", 0.72, 3.15, 2.4, 0.3, font_size=10, bold=True)
    _add_textbox(slide, _safe_str(bp.get("target_market"), "Target market pending"), 0.72, 3.42, 12.0, 0.6, font_size=16, bold=True)


def add_market_slide(pptx: Presentation, state: dict[str, Any], images: dict[str, bytes] | None = None):
    slide = pptx.slides.add_slide(pptx.slide_layouts[6])
    _configure_slide(slide)
    _add_image(slide, (images or {}).get("market"), 0.0, 0.0, 13.333, 7.5, opacity=0.08)
    _make_slide_title(slide, "Market Size", "TAM, SAM, and SOM presented as a nested funnel")
    market = state.get("market") or {}
    levels = [
        ("TAM", _safe_str(market.get("tam")), "#6C47FF"),
        ("SAM", _safe_str(market.get("sam")), "#00D4AA"),
        ("SOM", _safe_str(market.get("som")), "#4DA3FF"),
    ]
    widths = [7.8, 5.8, 3.8]
    for idx, ((name, value, accent), width) in enumerate(zip(levels, widths)):
        left = (12.0 - width) / 2
        top = 1.9 + idx * 1.05
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(0.82))
        shape.fill.background()
        shape.line.color.rgb = _rgb(accent)
        shape.line.width = Pt(1.7)
        _add_textbox(slide, f"{name}  •  {value}", left + 0.2, top + 0.16, width - 0.4, 0.25, font_size=18, bold=True, align=PP_ALIGN.CENTER)
        _add_textbox(slide, "Nested market segment", left + 0.2, top + 0.42, width - 0.4, 0.2, font_size=10, align=PP_ALIGN.CENTER)

    comps = (market.get("competitors") or [])[:3]
    x = 0.72
    for comp in comps:
        _add_card(slide, _safe_str(comp.get("name"), "Competitor"), _safe_str(comp.get("pricing"), "Pricing pending"), x, 5.0, 3.6, 1.15, accent="#2A2A35")
        x += 3.85


def add_architecture_slide(pptx: Presentation, state: dict[str, Any], images: dict[str, bytes] | None = None):
    slide = pptx.slides.add_slide(pptx.slide_layouts[6])
    _configure_slide(slide)
    _add_image(slide, (images or {}).get("architecture"), 0.0, 0.0, 13.333, 7.5, opacity=0.06)
    _make_slide_title(slide, "MVP Architecture", "Rendered directly from architecture.nodes and architecture.edges")
    architecture = state.get("mvp") or {}
    model = state.get("architecture") or state.get("mvp_architecture") or {}
    layers = model.get("layers") or []
    nodes = model.get("nodes") or []
    edges = model.get("edges") or []
    if not layers:
        layers = [
            {"id": "frontend", "label": "Frontend", "order": 1},
            {"id": "backend", "label": "Backend", "order": 2},
            {"id": "data", "label": "Data", "order": 3},
            {"id": "ai", "label": "AI", "order": 4},
        ]

    ordered_layers = sorted(layers, key=lambda x: x.get("order", 0))
    node_positions: dict[str, tuple[float, float, float, float]] = {}
    layer_height = 0.92
    start_y = 1.8
    for li, layer in enumerate(ordered_layers):
        _add_textbox(slide, _safe_str(layer.get("label"), layer.get("id", "Layer")).upper(), 0.78, start_y + li * 1.1, 1.2, 0.2, font_size=9, bold=True)
        layer_nodes = [n for n in nodes if n.get("layer") == layer.get("id")]
        if not layer_nodes and li == 0:
            layer_nodes = [{"id": "web", "label": "React", "type": "frontend", "layer": "frontend", "description": "User interface"}]
        if not layer_nodes:
            continue
        node_width = min(2.0, max(1.5, 10.3 / max(len(layer_nodes), 1)))
        gap = 0.25
        total = len(layer_nodes) * node_width + (len(layer_nodes) - 1) * gap
        start_x = (12.4 - total) / 2
        for ni, node in enumerate(layer_nodes):
            x = start_x + ni * (node_width + gap)
            y = start_y + li * 1.1
            shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(node_width), Inches(layer_height))
            shape.fill.background()
            shape.line.color.rgb = _rgb("#6C47FF" if li % 2 == 0 else "#00D4AA")
            shape.line.width = Pt(1.6)
            _add_textbox(slide, _safe_str(node.get("label"), "Node"), x + 0.12, y + 0.14, node_width - 0.24, 0.2, font_size=13, bold=True)
            _add_textbox(slide, _safe_str(node.get("description"), "Description pending"), x + 0.12, y + 0.42, node_width - 0.24, 0.35, font_size=9)
            node_positions[_safe_str(node.get("id"))] = (x + node_width / 2, y + layer_height, node_width, layer_height)

    for edge in edges:
        source = node_positions.get(_safe_str(edge.get("from")))
        target = node_positions.get(_safe_str(edge.get("to")))
        if not source or not target:
            continue
        x1, y1, _, _ = source
        x2, y2, _, _ = target
        line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        line.line.color.rgb = _rgb("#6C47FF")
        line.line.width = Pt(1.5)
        _add_textbox(slide, _safe_str(edge.get("label"), ""), min(x1, x2) + 0.05, (y1 + y2) / 2 - 0.12, abs(x2 - x1) or 1, 0.2, font_size=8, color="#00D4AA", align=PP_ALIGN.CENTER)


def add_tech_stack_slide(pptx: Presentation, state: dict[str, Any], images: dict[str, bytes] | None = None):
    slide = pptx.slides.add_slide(pptx.slide_layouts[6])
    _configure_slide(slide)
    _add_image(slide, (images or {}).get("tech_stack"), 0.0, 0.0, 13.333, 7.5, opacity=0.06)
    _make_slide_title(slide, "Recommended Tech Stack", "Layered implementation stack with complexity and rationale")
    stack = ((state.get("mvp") or {}).get("recommended_stack") or [])[:4]
    x = 0.72
    for item in stack:
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(1.9), Inches(2.95), Inches(1.65))
        shape.fill.solid()
        shape.fill.fore_color.rgb = _rgb("#111118")
        shape.line.color.rgb = _rgb("#2A2A35")
        shape.line.width = Pt(1.2)
        _add_tag(slide, _safe_str(item.get("layer"), "Layer"), x + 0.15, 2.02, 0.9, 0.22, fill="#0A0A0F")
        _add_textbox(slide, _safe_str(item.get("technology"), "Tech"), x + 0.15, 2.32, 2.55, 0.28, font_size=18, bold=True, color="#FFFFFF")
        _add_textbox(slide, _safe_str(item.get("reason"), "Reason pending"), x + 0.15, 2.62, 2.55, 0.46, font_size=10, color="#888899")
        _add_tag(slide, f"Complexity: {_safe_str(item.get('complexity'), 'Medium')}", x + 0.15, 3.17, 1.3, 0.24, fill="#111118", color="#00D4AA")
        x += 3.1


def add_financial_slide(pptx: Presentation, state: dict[str, Any], images: dict[str, bytes] | None = None):
    slide = pptx.slides.add_slide(pptx.slide_layouts[6])
    _configure_slide(slide)
    _add_image(slide, (images or {}).get("financial"), 0.0, 0.0, 13.333, 7.5, opacity=0.06)
    _make_slide_title(slide, "Financial Projection", "Five-year revenue, EBITDA, and free cash flow")
    financials = state.get("financials") or {}
    projections = financials.get("projections") or []
    data = CategoryChartData()
    data.categories = [f"Y{_safe_str(row.get('year'), '')}" for row in projections[:5]]
    if not projections:
        projections = [
            {"year": 1, "revenue": 120000, "ebitda": -780000, "fcf": -920000},
            {"year": 2, "revenue": 550000, "ebitda": -450000, "fcf": -510000},
            {"year": 3, "revenue": 1800000, "ebitda": 150000, "fcf": 90000},
            {"year": 4, "revenue": 4200000, "ebitda": 980000, "fcf": 760000},
            {"year": 5, "revenue": 8400000, "ebitda": 2200000, "fcf": 1850000},
        ]
        data.categories = [f"Y{row['year']}" for row in projections]
    data.add_series("Revenue", [float(row.get("revenue", 0)) for row in projections[:5]])
    data.add_series("EBITDA", [float(row.get("ebitda", 0)) for row in projections[:5]])
    data.add_series("FCF", [float(row.get("fcf", 0)) for row in projections[:5]])
    chart = slide.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(0.7), Inches(1.9), Inches(8.3), Inches(4.8), data).chart
    chart.has_legend = True
    chart.legend.include_in_layout = False
    chart.chart_title.has_text_frame = False
    plot = chart.plots[0]
    plot.has_data_labels = False
    _add_card(slide, "NPV", _safe_str(financials.get("npv"), "Pending"), 9.35, 1.95, 2.7, 1.0, accent="#6C47FF")
    _add_card(slide, "IRR", _safe_str(financials.get("irr"), "Pending"), 9.35, 3.1, 2.7, 1.0, accent="#00D4AA")
    _add_card(slide, "Payback", _safe_str(financials.get("payback_months"), "Pending"), 9.35, 4.25, 2.7, 1.0, accent="#4DA3FF")


def add_roadmap_slide(pptx: Presentation, state: dict[str, Any], images: dict[str, bytes] | None = None):
    slide = pptx.slides.add_slide(pptx.slide_layouts[6])
    _configure_slide(slide)
    _add_image(slide, (images or {}).get("roadmap"), 0.0, 0.0, 13.333, 7.5, opacity=0.06)
    _make_slide_title(slide, "Roadmap", "Milestones, week markers, deliverables, and task cards")
    roadmap = ((state.get("mvp") or {}).get("roadmap_phases") or (state.get("roadmapTimeline") or {}).get("phases") or [])
    if not roadmap:
        roadmap = [
            {"phase": 1, "title": "Foundation", "startWeek": 1, "endWeek": 2, "tasks": ["Authentication", "Landing Page", "Database"], "deliverables": ["Auth flow"]},
            {"phase": 2, "title": "Core Product", "startWeek": 3, "endWeek": 4, "tasks": ["Primary workflow", "API", "Dashboard"], "deliverables": ["Core flow"]},
            {"phase": 3, "title": "Launch", "startWeek": 5, "endWeek": 6, "tasks": ["Beta release", "Telemetry"], "deliverables": ["Public beta"]},
        ]
    x = 0.8
    y_line = 3.2
    timeline = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(y_line), Inches(11.4), Inches(0.04))
    timeline.fill.solid()
    timeline.fill.fore_color.rgb = _rgb("#2A2A35")
    timeline.line.fill.background()
    for idx, phase in enumerate(roadmap[:4]):
        milestone_x = x + idx * 3.15
        dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(milestone_x), Inches(y_line - 0.14), Inches(0.22), Inches(0.22))
        dot.fill.solid()
        dot.fill.fore_color.rgb = _rgb("#00D4AA" if idx == 0 else "#6C47FF")
        dot.line.fill.background()
        _add_textbox(slide, f"Week {_safe_str(phase.get('startWeek'), '?')}-{_safe_str(phase.get('endWeek'), '?')}", milestone_x - 0.25, 2.35, 2.2, 0.2, font_size=9, color="#888899", align=PP_ALIGN.CENTER)
        _add_textbox(slide, _safe_str(phase.get("title"), f"Phase {idx+1}"), milestone_x - 0.7, 2.56, 2.9, 0.22, font_size=13, bold=True, color="#FFFFFF", align=PP_ALIGN.CENTER)
        _add_tag(slide, _safe_str((phase.get("deliverables") or ["Milestone"])[0]), milestone_x - 0.35, 3.52, 2.0, 0.22, fill="#111118", color="#00D4AA")
        tasks = phase.get("tasks") or []
        for ti, task in enumerate(tasks[:3]):
            _add_tag(slide, _safe_str(task), milestone_x - 0.55, 3.86 + ti * 0.3, 2.4, 0.22, fill="#0A0A0F", color="#F0F0F0")


def add_competitor_slide(pptx: Presentation, state: dict[str, Any], images: dict[str, bytes] | None = None):
    slide = pptx.slides.add_slide(pptx.slide_layouts[6])
    _configure_slide(slide)
    _add_image(slide, (images or {}).get("competitor"), 0.0, 0.0, 13.333, 7.5, opacity=0.06)
    _make_slide_title(slide, "Competitor Landscape", "Key competitors and market positioning")
    competitors = (((state.get("market") or {}).get("competitors")) or [])[:4]
    headers = ["Company", "Founded", "Funding", "Pricing", "Focus", "Threat"]
    x = 0.72
    y = 1.95
    col_w = [2.0, 1.15, 1.25, 1.35, 3.05, 1.1]
    total = sum(col_w)
    # header row
    for i, h in enumerate(headers):
        left = x + sum(col_w[:i])
        _add_tag(slide, h, left, y, col_w[i] - 0.05, 0.24, fill="#111118", color="#888899")
    for r, comp in enumerate(competitors):
        row_y = y + 0.35 + r * 0.72
        values = [
            _safe_str(comp.get("name"), "Competitor"),
            _safe_str(comp.get("founded"), "—"),
            _safe_str(comp.get("funding"), "—"),
            _safe_str(comp.get("pricing"), "—"),
            _safe_str(comp.get("focus"), "—"),
            _safe_str(comp.get("threat_level"), "—"),
        ]
        for i, value in enumerate(values):
            left = x + sum(col_w[:i])
            _add_tag(slide, value, left, row_y, col_w[i] - 0.05, 0.28, fill="#0A0A0F", color="#F0F0F0")


def add_risks_slide(pptx: Presentation, state: dict[str, Any], images: dict[str, bytes] | None = None):
    slide = pptx.slides.add_slide(pptx.slide_layouts[6])
    _configure_slide(slide)
    _add_image(slide, (images or {}).get("risks"), 0.0, 0.0, 13.333, 7.5, opacity=0.06)
    _make_slide_title(slide, "Risks & Recommendations", "Operational risks plus investor-facing guidance")
    risks = (((state.get("business_plan") or {}).get("key_risks")) or ["Market timing", "Data quality", "Distribution"])[:3]
    recs = [
        "Launch with one tightly scoped workflow",
        "Instrument activation and retention from day one",
        "Keep the AI layer assistive before automation",
    ]
    _add_textbox(slide, "Risks", 0.72, 1.9, 2.2, 0.2, font_size=12, bold=True, color="#FF4D4F")
    for i, risk in enumerate(risks):
        _add_card(slide, f"Risk {i+1}", _safe_str(risk), 0.72 + i * 4.0, 2.15, 3.7, 0.95, accent="#FF4D4F")
    _add_textbox(slide, "Recommendations", 0.72, 3.6, 2.4, 0.2, font_size=12, bold=True, color="#00D4AA")
    for i, rec in enumerate(recs):
        _add_tag(slide, rec, 0.72, 3.9 + i * 0.42, 11.4, 0.28, fill="#111118", color="#F0F0F0")


import os

def generate_pitch_deck_pptx(state: dict[str, Any], images: dict[str, bytes] | None = None) -> BytesIO:
    template_path = os.path.join(os.path.dirname(__file__), "..", "..", "assets", "template.pptx")
    if os.path.exists(template_path):
        pptx = Presentation(template_path)
    else:
        pptx = Presentation()

    pptx.slide_width = SLIDE_W
    pptx.slide_height = SLIDE_H
    add_cover_slide(pptx, state, images)
    add_executive_summary_slide(pptx, state, images)
    add_market_slide(pptx, state, images)
    add_architecture_slide(pptx, state, images)
    add_tech_stack_slide(pptx, state, images)
    add_financial_slide(pptx, state, images)
    add_roadmap_slide(pptx, state, images)
    add_competitor_slide(pptx, state, images)
    add_risks_slide(pptx, state, images)
    output = BytesIO()
    pptx.save(output)
    output.seek(0)
    return output
