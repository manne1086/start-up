from __future__ import annotations

from io import BytesIO
from typing import Any

import httpx
from pptx import Presentation

from core.config import settings


def _safe_str(value: Any, fallback: str = "") -> str:
    if value is None:
        return fallback
    if isinstance(value, (int, float)):
        return f"{value:,}"
    text = str(value).strip()
    return text if text else fallback


def _join_list(values: Any, fallback: str = "") -> str:
    if not isinstance(values, list) or not values:
        return fallback
    cleaned = [_safe_str(value) for value in values if _safe_str(value)]
    return ", ".join(cleaned) if cleaned else fallback


def build_presentations_topic(state: dict[str, Any]) -> str:
    startup_name = _safe_str(state.get("startup_name"), "")
    idea = _safe_str(state.get("idea"), "A startup concept")
    business = state.get("business_plan") or {}
    market = state.get("market") or {}
    financials = state.get("financials") or {}
    mvp = state.get("mvp") or {}

    competitors = market.get("competitors") or []
    competitor_names = [item.get("name") for item in competitors[:3] if item.get("name")]

    topic_lines = [
        f"Create a professional investor pitch deck for a startup called '{startup_name or 'Our Startup'}'.",
        f"The startup idea: {idea}.",
        "",
        "This pitch deck is for pitching to investors and clients. Focus entirely on the startup business, NOT on the tool or platform used to create this deck.",
        "",
        "Slide 1 — Title: Show the startup name, tagline, and what it does in one sentence.",
        "",
        "Slide 2 — Problem:",
        f"  {_safe_str(business.get('problem_statement'), 'Describe the core problem this startup solves.')}",
        "",
        "Slide 3 — Solution:",
        f"  {_safe_str(business.get('solution'), 'Describe the solution.')}",
        f"  Value proposition: {_safe_str(business.get('value_proposition'), '')}",
        "",
        "Slide 4 — Target Market:",
        f"  {_safe_str(business.get('target_market'), 'Define the target audience.')}",
        "",
        "Slide 5 — Market Size:",
        f"  TAM: {_safe_str(market.get('tam'), 'TBD')}",
        f"  SAM: {_safe_str(market.get('sam'), 'TBD')}",
        f"  SOM: {_safe_str(market.get('som'), 'TBD')}",
    ]

    if competitor_names:
        topic_lines.append(f"  Key competitors: {', '.join(competitor_names)}")

    if market.get("market_gaps"):
        topic_lines.append(f"  Market gaps: {_join_list(market['market_gaps'])}")

    topic_lines += [
        "",
        "Slide 6 — Business Model & Revenue:",
        f"  Revenue model: {_safe_str(business.get('revenue_model'), 'TBD')}",
        f"  Pricing: {_safe_str(business.get('pricing'), 'TBD')}",
        f"  Go-to-market: {_safe_str(business.get('gtm_strategy'), 'TBD')}",
        "",
        "Slide 7 — Financial Projections:",
        f"  NPV: {_safe_str(financials.get('npv'), 'TBD')}",
        f"  IRR: {_safe_str(financials.get('irr'), 'TBD')}",
        f"  Payback period: {_safe_str(financials.get('payback_months'), 'TBD')} months",
    ]

    projections = financials.get("projections") or []
    if projections:
        topic_lines.append("  Revenue projections:")
        for proj in projections[:3]:
            topic_lines.append(f"    Year {proj.get('year', '?')}: Revenue {_safe_str(proj.get('revenue'), 'TBD')}, EBITDA {_safe_str(proj.get('ebitda'), 'TBD')}")

    topic_lines += [
        "",
        "Slide 8 — Roadmap:",
        f"  MVP timeline: {_safe_str(mvp.get('estimated_weeks'), 'TBD')} weeks",
        f"  Team size needed: {_safe_str(mvp.get('team_size'), 'TBD')}",
        f"  Estimated cost: {_safe_str(mvp.get('estimated_cost_inr'), 'TBD')}",
    ]

    roadmap = mvp.get("roadmap_phases") or []
    if roadmap:
        for phase in roadmap[:4]:
            topic_lines.append(f"    Phase {phase.get('phase', '?')}: {phase.get('title', '')} ({phase.get('weeks', '')})")

    risks = business.get("key_risks") or []
    mitigations = business.get("mitigation_steps") or []
    if risks:
        topic_lines += [
            "",
            "Slide 9 — Risks & Mitigation:",
        ]
        for i, risk in enumerate(risks[:4]):
            mitigation = mitigations[i] if i < len(mitigations) else ""
            topic_lines.append(f"  - Risk: {risk}" + (f" → Mitigation: {mitigation}" if mitigation else ""))

    topic_lines += [
        "",
        "Slide 10 — Call to Action:",
        f"  Invite investors to join {startup_name or 'this venture'}. Summarize the opportunity.",
        "",
        "IMPORTANT: This deck is about the startup described above. Do NOT mention VentureForge, AI generators, or any tool used to create this presentation.",
    ]

    return "\n".join(topic_lines)


def parse_pptx_slides(pptx_bytes: bytes) -> list[dict[str, Any]]:
    """Parse a PPTX file and extract slide content for the in-app editor."""
    prs = Presentation(BytesIO(pptx_bytes))
    slides: list[dict[str, Any]] = []

    for i, slide in enumerate(prs.slides):
        title = ""
        body_parts: list[str] = []
        bullets: list[str] = []

        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue

            is_title = False
            if hasattr(shape, "placeholder_format") and shape.placeholder_format:
                if shape.placeholder_format.idx in (0, 15):
                    is_title = True
            if not is_title and slide.shapes.title and shape.shape_id == slide.shapes.title.shape_id:
                is_title = True

            if is_title:
                title = shape.text_frame.text.strip()
            else:
                for para in shape.text_frame.paragraphs:
                    text = para.text.strip()
                    if not text:
                        continue
                    if para.level and para.level > 0:
                        bullets.append(text)
                    else:
                        if len(text) < 120:
                            bullets.append(text)
                        else:
                            body_parts.append(text)

        if not title and body_parts:
            title = body_parts.pop(0)

        content: dict[str, Any] = {}
        if body_parts:
            content["text"] = " ".join(body_parts)
        if bullets:
            content["bullets"] = bullets

        slide_type = "title" if i == 0 else "content"
        title_lower = title.lower()
        if "market" in title_lower and ("tam" in title_lower or "sam" in title_lower or "size" in title_lower):
            slide_type = "market"

        slides.append({
            "number": i + 1,
            "title": title or f"Slide {i + 1}",
            "type": slide_type,
            "content": content,
        })

    return slides


def _build_payload(state: dict[str, Any]) -> dict[str, Any]:
    return {
        "content": build_presentations_topic(state),
        "n_slides": 10,
        "language": "English",
        "template": "general",
        "export_as": "pptx",
    }


def _get_headers() -> dict[str, str]:
    return {
        "Authorization": f"Bearer {settings.PRESENTATIONS_AI_API_KEY}",
        "Content-Type": "application/json",
    }


def _resolve_url(url: str) -> str:
    if url.startswith("/"):
        return settings.PRESENTATIONS_AI_BASE_URL.rstrip("/") + url
    return url


async def generate_presenton_presentation(state: dict[str, Any]) -> dict[str, str]:
    """Call Presenton API and return metadata (presentation_id, path, edit_path).
    Used during the agent pipeline to pre-generate the PPTX."""
    if not settings.PRESENTATIONS_AI_API_KEY:
        raise RuntimeError("Presenton API key is not configured")

    payload = _build_payload(state)
    headers = _get_headers()

    async with httpx.AsyncClient(base_url=settings.PRESENTATIONS_AI_BASE_URL, timeout=180) as client:
        response = await client.post("/api/v1/ppt/presentation/generate", headers=headers, json=payload)
        if response.status_code != 200:
            try:
                detail = response.json()
            except Exception:
                detail = response.text
            raise RuntimeError(f"Presenton API {response.status_code}: {detail}")
        data = response.json()

    return {
        "presentation_id": data.get("presentation_id", ""),
        "path": _resolve_url(data.get("path", "")),
        "edit_path": _resolve_url(data.get("edit_path", "")),
    }


def generate_presentations_ai_pptx(state: dict[str, Any]) -> BytesIO:
    """Synchronous version: call Presenton and return the PPTX bytes.
    Used as fallback in the export route."""
    if not settings.PRESENTATIONS_AI_API_KEY:
        raise RuntimeError("Presenton API key is not configured")

    payload = _build_payload(state)
    headers = _get_headers()

    with httpx.Client(base_url=settings.PRESENTATIONS_AI_BASE_URL, timeout=180) as client:
        response = client.post("/api/v1/ppt/presentation/generate", headers=headers, json=payload)
        if response.status_code != 200:
            try:
                detail = response.json()
            except Exception:
                detail = response.text
            raise RuntimeError(f"Presenton API {response.status_code}: {detail}")
        data = response.json()

    download_url = _resolve_url(data.get("path", ""))
    if not download_url:
        raise RuntimeError(f"Presenton did not return a download URL: {data}")

    file_response = httpx.get(download_url, timeout=120)
    file_response.raise_for_status()
    return BytesIO(file_response.content)
